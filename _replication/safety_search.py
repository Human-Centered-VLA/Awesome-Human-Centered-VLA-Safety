#!/usr/bin/env python3
"""Crawl, rebuild, and publish the seven VLA Research spreadsheet tabs.

Paper classifications and evidence are explicit reviewed inputs.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import time
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "examples" / "vla_research_reference.json"
EXAMPLE_CRAWL = ROOT / "examples" / "crawl_snapshot.csv"
API_URL = "https://export.arxiv.org/api/query"
ATOM = {"a": "http://www.w3.org/2005/Atom"}


def schema() -> dict:
    return json.loads((ROOT / "sheet_schema.json").read_text(encoding="utf-8"))


def clean(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime) and value.time() == dt.time():
        return value.date().isoformat()
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def write_csv(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)


def read_table(path: Path, headers: list[str]) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        if next(reader, None) != headers:
            raise ValueError(f"Header mismatch in {path}")
        rows = list(reader)
        if any(len(row) != len(headers) for row in rows):
            raise ValueError(f"Row width mismatch in {path}")
        return rows


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def table_digest(rows: list[list[str]]) -> str:
    return hashlib.sha256(json.dumps(rows, ensure_ascii=False, separators=(",", ":")).encode("utf-8")).hexdigest()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def parse_atom(xml_text: str, query: str) -> list[list[str]]:
    """Map an arXiv Atom page to the exact Raw Papers column order."""
    rows = []
    for entry in ET.fromstring(xml_text).findall("a:entry", ATOM):
        entry_id = entry.findtext("a:id", default="", namespaces=ATOM).rstrip("/")
        paper_id = entry_id.rsplit("/", 1)[-1]
        if not re.fullmatch(r"\d{4}\.\d{4,5}(?:v\d+)?", paper_id):
            continue
        title = clean(entry.findtext("a:title", default="", namespaces=ATOM))
        abstract = clean(entry.findtext("a:summary", default="", namespaces=ATOM))
        authors = "; ".join(
            clean(a.findtext("a:name", default="", namespaces=ATOM))
            for a in entry.findall("a:author", ATOM)
        )
        categories = "; ".join(c.get("term", "") for c in entry.findall("a:category", ATOM))
        rows.append([
            paper_id, title[:500], authors[:1000], abstract[:1000],
            entry.findtext("a:published", default="", namespaces=ATOM)[:10],
            entry.findtext("a:updated", default="", namespaces=ATOM)[:10],
            categories[:300], "example", query[:120], "", "", "", "", "",
            "", f"Example crawl query: {query}"[:500], "", "", "",
        ])
    return rows


def crawl(args: argparse.Namespace) -> None:
    import requests

    query = args.query
    if args.start_date or args.end_date:
        if not (args.start_date and args.end_date):
            raise ValueError("Provide both --start-date and --end-date")
        if args.start_date > args.end_date:
            raise ValueError("--start-date must be on or before --end-date")
        query += f" AND submittedDate:[{args.start_date}0000 TO {args.end_date}2359]"

    started = dt.datetime.now(dt.timezone.utc).isoformat()
    rows: list[list[str]] = []
    seen: set[str] = set()
    pages = []
    with requests.Session() as session:
        for start in range(0, args.max_results, args.page_size):
            params = {
                "search_query": query, "start": start,
                "max_results": min(args.page_size, args.max_results - start),
                "sortBy": "submittedDate", "sortOrder": "ascending",
            }
            for attempt in range(3):
                response = session.get(
                    API_URL, params=params,
                    headers={"User-Agent": "vla-research-example/1.0"},
                    timeout=args.timeout,
                )
                if response.status_code not in {429, 500, 502, 503, 504}:
                    response.raise_for_status()
                    break
                if attempt == 2:
                    response.raise_for_status()
                retry_after = response.headers.get("Retry-After", "")
                time.sleep(float(retry_after) if retry_after.isdigit() else 3 * 2**attempt)
            page = parse_atom(response.text, args.query)
            pages.append({"params": params, "response_sha256": hashlib.sha256(response.content).hexdigest()})
            for row in page:
                if row[0] not in seen:
                    seen.add(row[0])
                    rows.append(row)
            if len(page) < params["max_results"]:
                break
            if start + args.page_size < args.max_results:
                time.sleep(args.sleep)
    write_csv(args.output, schema()["sheets"]["Raw Papers"], rows)
    write_json(args.output.with_suffix(".manifest.json"), {
        "started_utc": started, "api_url": API_URL, "query": query,
        "max_results": args.max_results, "page_size": args.page_size,
        "pages": pages, "rows": len(rows), "csv_sha256": digest(args.output),
        "note": "Scores and full-text evidence are blank until reviewed; API results can change.",
    })
    print(f"Wrote {len(rows)} example papers to {args.output}")


def read_reference(path: Path) -> dict[str, list[list[str]]]:
    reference_schema = schema()
    expected = reference_schema["sheets"]
    if path.suffix.lower() == ".json":
        values = json.loads(path.read_text(encoding="utf-8"))["sheets"]
        if list(values) != list(expected):
            raise ValueError("Reference snapshot must contain the seven tabs in schema order")
        for title, headers in expected.items():
            if not values[title] or values[title][0] != headers:
                raise ValueError(f"Reference header mismatch in {title}")
            if any(len(row) != len(headers) or any(not isinstance(v, str) for v in row) for row in values[title]):
                raise ValueError(f"Reference row mismatch in {title}")
        return {title: rows[1:] for title, rows in values.items()}
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("Install dependencies: pip install -r requirements.txt") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    tabs: dict[str, list[list[str]]] = {}
    try:
        if len(workbook.worksheets) != len(expected):
            raise ValueError("Reference workbook must contain seven tabs")
        for (title, headers), sheet in zip(expected.items(), workbook.worksheets):
            if sheet.title != reference_schema["excel_titles"][title]:
                raise ValueError(f"Reference tab mismatch: expected {title}, found {sheet.title}")
            values = [[decode_excel(cell_text(v)) for v in row] for row in sheet.iter_rows(values_only=True)]
            while values and not any(values[-1]):
                values.pop()
            if not values or values[0][:len(headers)] != headers:
                raise ValueError(f"Reference header mismatch in {title}")
            if any(any(row[len(headers):]) for row in values):
                raise ValueError(f"Unexpected populated columns in {title}")
            tabs[title] = [(row + [""] * len(headers))[:len(headers)] for row in values[1:]]
        return tabs
    finally:
        workbook.close()


def encode_excel(value: str) -> str:
    """Preserve literal escape tokens and text characters forbidden in XML."""
    value = re.sub(r"_x[0-9A-Fa-f]{4}_", lambda m: "_x005F_" + m.group()[1:], value)
    return re.sub(r"[\x00-\x08\x0b-\x0d\x0e-\x1f\ufffe\uffff]",
                  lambda m: f"_x{ord(m.group()):04X}_", value)


def decode_excel(value: str) -> str:
    return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), value)


def google_value(title: str, column: str, value: str) -> dict:
    """Keep native dates and numeric columns; all other values are literal text."""
    if not value:
        return {}
    if title == "Raw Papers" and column in {"published", "updated"}:
        try:
            date = dt.date.fromisoformat(value)
            return {"numberValue": (date - dt.date(1899, 12, 30)).days}
        except ValueError:
            pass
    if column in {"No.", "Year", "source_vla_hits", "source_embodiment_hits",
                  "source_safety_hits", "full_text_safety_hits", "metadata_score", "full_text_score"}:
        try:
            number = float(value)
            if number == number and abs(number) != float("inf"):
                return {"numberValue": number}
        except ValueError:
            pass
    return {"stringValue": value}


def rebuild_workbook(path: Path, tabs: dict[str, list[list[str]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers in schema()["sheets"].items():
        sheet = workbook.create_sheet(schema()["excel_titles"][title])
        for row_index, row in enumerate([headers] + tabs[title], 1):
            for column_index, value in enumerate(row, 1):
                typed = google_value(title, headers[column_index - 1], value) if row_index > 1 else {"stringValue": value}
                stored = next(iter(typed.values()), None)
                cell = sheet.cell(row_index, column_index, encode_excel(stored) if isinstance(stored, str) else stored)
                if "stringValue" in typed:
                    cell.data_type = "s"  # Evidence starting with '=' must stay literal text.
                if title == "Raw Papers" and headers[column_index - 1] in {"published", "updated"} and "numberValue" in typed:
                    cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="334155")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 24 if len(header) < 20 else 44
    workbook.save(path)
    workbook.close()


def analyze(args: argparse.Namespace) -> None:
    expected = schema()["sheets"]
    reference = read_reference(args.reference)
    tabs = {title: [row[:] for row in rows] for title, rows in reference.items()}
    reviewed_dir = getattr(args, "reviewed_dir", None)
    if reviewed_dir:
        for title, headers in expected.items():
            if title != "Raw Papers":
                tabs[title] = read_table(reviewed_dir / f"{title}.csv", headers)
    crawl_rows = read_table(args.crawl, expected["Raw Papers"])
    if getattr(args, "use_crawl", False):
        tabs["Raw Papers"] = crawl_rows
        ids = {row[0] for row in crawl_rows}
        for title, headers in expected.items():
            if "ArXiv ID" in headers:
                index = headers.index("ArXiv ID")
                tabs[title] = [row for row in tabs[title] if row[index] in ids]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for title, headers in expected.items():
        write_csv(args.output_dir / "tabs" / f"{title}.csv", headers, tabs[title])
    rebuilt = args.output_dir / "VLA Research.xlsx"
    rebuild_workbook(rebuilt, tabs)
    if read_reference(rebuilt) != tabs:
        raise ValueError("Rebuilt workbook failed its cell-by-cell round trip")

    reference_ids = [r[0] for r in reference["Raw Papers"] if r[0]]
    crawl_ids = [r[0] for r in crawl_rows if r[0]]
    reference_set, crawl_set = set(reference_ids), set(crawl_ids)
    def tab_ids(title: str, column: int) -> set[str]:
        return {row[column] for row in tabs[title]
                if len(row) > column and row[column]}

    master = tab_ids("Sorted Master Papers", 0)
    landscape = tab_ids("Sorted Master Papers _ landscape+note", 1)
    technical = tab_ids("Sorted Master Papers _ technical", 1)
    borderline = tab_ids("Sorted Borderline Papers", 0)
    report = {
        "source_spreadsheet": schema()["source_spreadsheet"],
        "mode": "crawl_with_reviewed_inputs" if getattr(args, "use_crawl", False) else "reference_reproduction",
        "reference_rows_by_tab": {title: sum(any(row) for row in rows) for title, rows in reference.items()},
        "output_rows_by_tab": {title: sum(any(row) for row in rows) for title, rows in tabs.items()},
        "crawl_rows": len(crawl_rows),
        "reference_duplicate_ids": sorted(
            paper_id for paper_id, n in Counter(reference_ids).items() if n > 1
        ),
        "crawl_only_ids": sorted(crawl_set - reference_set),
        "reference_only_ids": sorted(reference_set - crawl_set),
        "curation_checks": {
            "landscape_not_in_master": sorted(landscape - master),
            "technical_not_in_master": sorted(technical - master),
            "borderline_also_in_master": sorted(borderline & master),
        },
        "unclassified_output_ids": sorted(tab_ids("Raw Papers", 0) - master - borderline),
        "note": "Crawl-only papers require human review before entering curated tabs.",
    }
    write_json(args.output_dir / "alignment_report.json", report)
    write_json(args.output_dir / "manifest.json", {
        "mode": report["mode"], "source_spreadsheet": schema()["source_spreadsheet"],
        "inputs": {"reference_sha256": digest(args.reference), "crawl_sha256": digest(args.crawl),
                   "schema_sha256": digest(ROOT / "sheet_schema.json"),
                   "script_sha256": digest(Path(__file__)),
                   "reviewed_csv_sha256": {title: digest(reviewed_dir / f"{title}.csv")
                                           for title in expected if reviewed_dir and title != "Raw Papers"}},
        "tabs": {title: {"columns": len(headers), "data_rows": len(tabs[title]),
                         "content_sha256": table_digest([headers] + tabs[title]),
                         "csv_sha256": digest(args.output_dir / "tabs" / f"{title}.csv")}
                 for title, headers in expected.items()},
        "workbook_round_trip": "all cells match",
        "analysis_method": "Rebuild from explicit reviewed classifications, evidence, and milestones.",
    })
    print(f"Rebuilt and verified seven tabs in {args.output_dir}; {len(report['crawl_only_ids'])} crawl-only IDs need review.")


def yyyymmdd(value: str) -> str:
    try:
        dt.datetime.strptime(value, "%Y%m%d")
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Use YYYYMMDD") from exc
    return value


def publish(args: argparse.Namespace) -> None:
    """Copy the native template, apply changed values, and verify every output cell."""
    import requests

    token = os.environ.get("GOOGLE_ACCESS_TOKEN")
    if not token:
        raise ValueError("Set GOOGLE_ACCESS_TOKEN to an OAuth access token with Drive access")
    expected = schema()["sheets"]
    tables = {title: [headers] + read_table(args.output_dir / "tabs" / f"{title}.csv", headers)
              for title, headers in expected.items()}
    manifest = json.loads((args.output_dir / "manifest.json").read_text(encoding="utf-8"))
    if manifest["inputs"]["schema_sha256"] != digest(ROOT / "sheet_schema.json"):
        raise ValueError("Schema changed; rerun analyze before publishing")
    for title in expected:
        if manifest["tabs"][title]["csv_sha256"] != digest(args.output_dir / "tabs" / f"{title}.csv"):
            raise ValueError(f"{title} changed; rerun analyze before publishing")
    source_id = schema()["source_spreadsheet"].split("/d/", 1)[1].split("/", 1)[0]
    drive = "https://www.googleapis.com/drive/v3/files"
    sheets = "https://sheets.googleapis.com/v4/spreadsheets"
    receipt = {"status": "starting", "source_id": source_id,
               "manifest_sha256": digest(args.output_dir / "manifest.json")}
    receipt_path = args.output_dir / "published_version.json"
    with requests.Session() as session:
        session.headers["Authorization"] = f"Bearer {token}"

        def api(method: str, url: str, **kwargs) -> dict:
            response = session.request(method, url, timeout=90, **kwargs)
            response.raise_for_status()
            return response.json()

        source = api("GET", f"{sheets}/{source_id}", params={"fields": "sheets.properties"})
        if [s["properties"]["title"] for s in source["sheets"]] != list(expected):
            raise ValueError("Live template tabs differ from the recorded schema")
        folder_id = args.folder_id
        if not folder_id:
            folders = api("GET", drive, params={
                "q": "name = 'ChatGPT' and mimeType = 'application/vnd.google-apps.folder' and 'root' in parents and trashed = false",
                "fields": "files(id)", "pageSize": 100,
            })["files"]
            folder_id = folders[0]["id"] if folders else api("POST", drive, json={
                "name": "ChatGPT", "mimeType": "application/vnd.google-apps.folder", "parents": ["root"],
            })["id"]
        copied = api("POST", f"{drive}/{source_id}/copy", params={"fields": "id"},
                     json={"name": args.title, "parents": [folder_id]})
        target_id = copied["id"]
        if target_id == source_id:
            raise ValueError("Destination must differ from the reference spreadsheet")
        receipt.update({"spreadsheet_id": target_id,
                        "url": f"https://docs.google.com/spreadsheets/d/{target_id}/edit", "status": "copied"})
        write_json(receipt_path, receipt)
        try:
            metadata = api("GET", f"{sheets}/{target_id}", params={"fields": "sheets.properties"})
            properties = {s["properties"]["title"]: s["properties"] for s in metadata["sheets"]}
            requests_to_send = []
            from openpyxl.utils import get_column_letter

            for title, table in tables.items():
                prop = properties[title]
                grid = prop["gridProperties"]
                if len(table) > grid["rowCount"]:
                    requests_to_send.append({"updateSheetProperties": {
                        "properties": {"sheetId": prop["sheetId"], "gridProperties": {"rowCount": len(table)}},
                        "fields": "gridProperties.rowCount",
                    }})
                    for paste_type in ("PASTE_FORMAT", "PASTE_DATA_VALIDATION"):
                        requests_to_send.append({"copyPaste": {
                            "source": {"sheetId": prop["sheetId"], "startRowIndex": grid["rowCount"] - 1,
                                       "endRowIndex": grid["rowCount"], "startColumnIndex": 0, "endColumnIndex": len(expected[title])},
                            "destination": {"sheetId": prop["sheetId"], "startRowIndex": grid["rowCount"],
                                            "endRowIndex": len(table), "startColumnIndex": 0, "endColumnIndex": len(expected[title])},
                            "pasteType": paste_type,
                        }})
                    grid["rowCount"] = len(table)
            if requests_to_send:
                api("POST", f"{sheets}/{target_id}:batchUpdate", json={"requests": requests_to_send})
            ranges = [f"'{title.replace(chr(39), chr(39)*2)}'!A1:{get_column_letter(len(expected[title]))}{properties[title]['gridProperties']['rowCount']}"
                      for title in tables]

            def read_values() -> list[dict]:
                return api("GET", f"{sheets}/{target_id}/values:batchGet", params={
                    "ranges": ranges, "valueRenderOption": "FORMATTED_VALUE",
                })["valueRanges"]

            def normalized(rows: list, width: int) -> list[list[str]]:
                result = [[cell_text(v) for v in row] + [""] * (width - len(row)) for row in rows]
                while result and not any(result[-1]):
                    result.pop()
                return result

            changes = []
            for (title, table), live in zip(tables.items(), read_values()):
                width = len(expected[title])
                current = normalized(live.get("values", []), width)
                if not current or current[0] != expected[title]:
                    raise ValueError(f"Live header mismatch in {title}")
                for row_index in range(max(len(table), len(current))):
                    old = current[row_index] if row_index < len(current) else [""] * width
                    new = table[row_index] if row_index < len(table) else [""] * width
                    for column_index, (before, after) in enumerate(zip(old, new)):
                        if before != after:
                            changes.append({"updateCells": {
                                "start": {"sheetId": properties[title]["sheetId"], "rowIndex": row_index, "columnIndex": column_index},
                                "rows": [{"values": [{"userEnteredValue": google_value(title, expected[title][column_index], after)}]}],
                                "fields": "userEnteredValue",
                            }})
            # Only changed literal cells are overwritten; matching native structures stay intact.
            if changes:
                cells = api("GET", f"{sheets}/{target_id}", params={
                    "ranges": ranges, "includeGridData": "true",
                    "fields": "sheets(properties(sheetId),data(startRow,startColumn,rowData(values(userEnteredValue,chipRuns))))",
                })
                protected = set()
                for sheet in cells["sheets"]:
                    for block in sheet.get("data", []):
                        for row_index, row in enumerate(block.get("rowData", []), block.get("startRow", 0)):
                            for col_index, cell in enumerate(row.get("values", []), block.get("startColumn", 0)):
                                if cell.get("chipRuns") or "formulaValue" in cell.get("userEnteredValue", {}):
                                    protected.add((sheet["properties"]["sheetId"], row_index, col_index))
                for change in changes:
                    start = change["updateCells"]["start"]
                    if (start["sheetId"], start["rowIndex"], start["columnIndex"]) in protected:
                        raise ValueError("A changed cell contains a formula or chip; review the copied version")
            changes.extend({"updateCells": {
                "start": {"sheetId": properties[title]["sheetId"], "rowIndex": 0, "columnIndex": 0},
                "rows": [{"values": [{"note": f"Reproduced with safety_search.py; manifest SHA-256: {receipt['manifest_sha256']}"}]}],
                "fields": "note",
            }} for title in tables)
            for start in range(0, len(changes), 100):
                api("POST", f"{sheets}/{target_id}:batchUpdate", json={"requests": changes[start:start + 100]})
            for (title, table), live in zip(tables.items(), read_values()):
                if normalized(live.get("values", []), len(expected[title])) != normalized(table, len(expected[title])):
                    raise ValueError(f"Published cell values differ in {title}")
            receipt.update({"status": "verified", "tabs": list(tables),
                            "verified_cells": sum(len(table) * len(expected[title]) for title, table in tables.items())})
            write_json(receipt_path, receipt)
        except Exception:
            receipt["status"] = "incomplete"
            write_json(receipt_path, receipt)
            print(f"Incomplete copy retained for inspection: {receipt['url']}")
            raise
    print(f"Created and verified: {receipt['url']}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    crawl_parser = commands.add_parser("crawl", help="Run a small example arXiv crawl")
    crawl_parser.add_argument("--query", required=True, help="arXiv API search expression")
    crawl_parser.add_argument("--start-date", type=yyyymmdd)
    crawl_parser.add_argument("--end-date", type=yyyymmdd)
    crawl_parser.add_argument("--max-results", type=int, default=20)
    crawl_parser.add_argument("--page-size", type=int, default=20)
    crawl_parser.add_argument("--sleep", type=float, default=3.0)
    crawl_parser.add_argument("--timeout", type=int, default=30)
    crawl_parser.add_argument("--output", type=Path, default=Path("output/crawl.csv"))
    analyze_parser = commands.add_parser(
        "analyze", help="Rebuild and verify all seven tabs from reviewed inputs"
    )
    analyze_parser.add_argument("--reference", type=Path, default=REFERENCE)
    analyze_parser.add_argument("--crawl", type=Path, default=EXAMPLE_CRAWL)
    analyze_parser.add_argument("--use-crawl", action="store_true", help="Use the crawl as Raw Papers and retain matching reviewed analyses")
    analyze_parser.add_argument("--reviewed-dir", type=Path, help="Six reviewed CSVs with exact tab names and headers")
    analyze_parser.add_argument("--output-dir", type=Path, default=Path("output"))
    publish_parser = commands.add_parser("publish", help="Create and verify a separate native Google Sheets version")
    publish_parser.add_argument("--output-dir", type=Path, default=Path("output"))
    publish_parser.add_argument("--title", default=f"VLA Research — reproduced {dt.date.today().isoformat()}")
    publish_parser.add_argument("--folder-id", help="Destination Drive folder; defaults to ChatGPT in My Drive")
    args = parser.parse_args()
    if args.command == "crawl":
        if args.max_results < 1 or not 1 <= args.page_size <= 100:
            parser.error("--max-results must be positive and --page-size must be 1..100")
        if args.sleep < 0 or args.timeout <= 0:
            parser.error("--sleep must be nonnegative and --timeout positive")
        crawl(args)
    elif args.command == "analyze":
        for flag, path in (("--reference", args.reference), ("--crawl", args.crawl)):
            if not path.is_file():
                parser.error(f"{flag} file not found: {path}; provide a local input file")
        analyze(args)
    else:
        publish(args)


if __name__ == "__main__":
    main()
